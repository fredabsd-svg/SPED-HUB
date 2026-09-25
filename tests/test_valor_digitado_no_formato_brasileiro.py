"""Valor digitado com vírgula decimal é número — e o que não é número é recusado.

O defeito: `desserializar` recebia "190,00" para uma coluna `Float`, não
conseguia converter, registrava um aviso no log e devolvia o **texto**. O
ajuste era gravado assim; o recálculo do cabeçalho tratava o texto como zero
(o C100 ia para 0,00) e a geração quebrava depois, com `TypeError` na EFD
ICMS/IPI e `decimal.InvalidOperation` na EFD-Contribuições — que a CLI nem
traduzia em mensagem. O erro de digitação aparecia no fechamento, longe de
quem digitou.

A regra agora: o formato brasileiro ("1.234,56", "190,00") é aceito, o
formato com ponto ("1234.56") continua valendo, e o resto é recusado com
`ValueError` **na hora de gravar** — no terminal, na tela, na planilha e na
ação de regra —, nunca depois.
"""

from __future__ import annotations

import datetime
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from src.cli import main
from src.db.models import (
    AjusteFiscal,
    DocumentoFiscal,
    Empresa,
    Escritorio,
    criar_engine,
    get_session,
    init_db,
)
from src.documentos import (
    ORIGEM_USUARIO,
    MotorDeClassificacao,
    RegraInvalida,
    Selecao,
    aplicar_ajuste,
    exportar,
    reimportar,
    valor_efetivo,
    valor_tipado,
)
from src.documentos.classificacao import criar_regra
from src.escrituracoes import GeradorEFDContribuicoes, GeradorEFDICMS
from src.escrituracoes.leiaute import EFD_CONTRIBUICOES, EFD_ICMS
from tests.fixtures_nfe import nfe_xml

CNPJ = "98765432000198"
INICIO = datetime.date(2026, 7, 1)
FIM = datetime.date(2026, 7, 31)


@pytest.fixture
def banco(tmp_path):
    """Uma entrada de um item (ICMS 180,00, PIS 16,50) e a URL do banco."""
    url = f"sqlite:///{tmp_path / 'fiscal.db'}"
    engine = criar_engine(url=url)
    init_db(engine)
    with get_session(engine) as sessao:
        sessao.add(Escritorio(nome="Teste", slug="teste"))
        sessao.commit()
        sessao.add(
            Empresa(
                cnpj=CNPJ,
                nome="COMERCIO EXEMPLO LTDA",
                uf="TO",
                ie="293456789",
                cod_mun="1721000",
                ind_perfil="A",
                ind_ativ="1",
                ind_ativ_contribuicoes="2",
                cod_inc_trib="1",
                escritorio_id=1,
            )
        )
        sessao.commit()
    engine.dispose()
    xml = tmp_path / "nfe.xml"
    xml.write_bytes(nfe_xml())
    assert main(["fiscal", "importar", str(xml), "--escritorio", "1", "--db", url]) == 0
    return url


def _alterar(banco, campo, valor, *extras):
    return main(
        [
            "fiscal",
            "alterar",
            "--empresa",
            "1",
            "--campo",
            campo,
            "--valor",
            valor,
            "--db",
            banco,
            *extras,
        ]
    )


def _gerar(banco, destino, tipo="efd_icms"):
    return main(
        [
            "fiscal",
            "gerar",
            "--empresa",
            "1",
            "--de",
            "2026-07-01",
            "--ate",
            "2026-07-31",
            "--tipo",
            tipo,
            "--saida",
            str(destino),
            "--db",
            banco,
        ]
    )


def _campo(arquivo, tipo: str, nome: str, leiaute=EFD_ICMS) -> str:
    """O campo pelo NOME, no primeiro registro do tipo — a posição vem do leiaute."""
    for linha in arquivo.read_bytes().decode("latin-1").split("\r\n"):
        if linha.startswith(f"|{tipo}|"):
            return linha.split("|")[2:-1][leiaute[tipo].index(nome)]
    raise AssertionError(f"{tipo} não está no arquivo")


class TestPelaLinhaDeComando:
    def test_virgula_decimal_chega_ao_arquivo(self, banco, tmp_path):
        assert _alterar(banco, "valor_icms", "190,00", "--confirmar") == 0
        destino = tmp_path / "efd.txt"

        assert _gerar(banco, destino) == 0, (
            "o valor digitado com vírgula foi aceito e a geração quebrou depois: "
            "o erro aparece no fechamento, longe de quem digitou"
        )
        assert _campo(destino, "C170", "VL_ICMS") == "190,00"
        assert (
            _campo(destino, "E110", "VL_TOT_CREDITOS") == "190,00"
        ), "o crédito do E110 não é o ICMS corrigido: o arquivo apura outro imposto"

    def test_o_cabecalho_e_recomposto_com_o_valor_digitado(self, banco, tmp_path):
        """O recálculo lia o texto como zero e zerava o VL_ICMS do C100."""
        _alterar(banco, "valor_icms", "190,00", "--confirmar")
        destino = tmp_path / "efd.txt"
        _gerar(banco, destino)

        assert _campo(destino, "C100", "VL_ICMS") == "190,00", (
            "o C100 saiu com o ICMS do cabeçalho zerado pelo recálculo, e não bate "
            "com a soma dos itens — é o que o validador confere"
        )

    def test_milhar_com_ponto_e_decimal_com_virgula(self, banco, tmp_path):
        assert _alterar(banco, "valor_pis", "1.234,56", "--confirmar") == 0
        destino = tmp_path / "contrib.txt"

        assert _gerar(banco, destino, tipo="efd_contribuicoes") == 0, (
            "a EFD-Contribuições quebrou com decimal.InvalidOperation num valor "
            "que a alteração em massa aceitou"
        )
        assert _campo(destino, "C170", "VL_PIS", EFD_CONTRIBUICOES) == "1234,56"

    def test_o_que_nao_e_numero_e_recusado_antes_de_gravar(self, banco, capsys):
        codigo = _alterar(banco, "valor_icms", "cento e noventa", "--confirmar")
        saida = capsys.readouterr().out

        assert codigo == 1
        assert (
            "valor_icms" in saida and "cento e noventa" in saida
        ), "a recusa não diz qual campo nem qual valor — quem digitou não sabe o que corrigir"
        with get_session(criar_engine(url=banco)) as sessao:
            assert (
                not sessao.execute(select(AjusteFiscal)).scalars().all()
            ), "o valor inválido foi gravado; a geração quebraria no fechamento"


class TestConversao:
    @pytest.mark.parametrize(
        ("bruto", "esperado"),
        [
            ("190,00", 190.0),
            ("1.234,56", 1234.56),
            ("1.234.567,8", 1234567.8),
            ("1234.56", 1234.56),
            ("1000", 1000.0),
            (" 12,5 ", 12.5),
            ("-3,25", -3.25),
        ],
    )
    def test_formatos_aceitos(self, bruto, esperado):
        assert valor_tipado("valor_icms", bruto) == pytest.approx(esperado)

    @pytest.mark.parametrize("bruto", ["abc", "1,2,3", "12,34.5", "1.23,4", "nan", "inf", ""])
    def test_o_resto_e_recusado(self, bruto):
        with pytest.raises(ValueError, match="valor_icms"):
            valor_tipado("valor_icms", bruto)

    def test_coluna_inteira_nao_aceita_fracao(self):
        assert valor_tipado("numero_item", "3") == 3
        with pytest.raises(ValueError, match="numero_item"):
            valor_tipado("numero_item", "3,5")


@pytest.fixture
def sessao(tmp_path):
    engine = criar_engine(url=f"sqlite:///{tmp_path / 'camada.db'}")
    init_db(engine)
    with get_session(engine) as s:
        escritorio = Escritorio(nome="Teste", slug="teste")
        s.add(escritorio)
        s.commit()
        s.add(
            Empresa(
                cnpj=CNPJ,
                nome="COMERCIO EXEMPLO LTDA",
                uf="TO",
                ie="293456789",
                cod_mun="1721000",
                ind_perfil="A",
                ind_ativ="1",
                ind_ativ_contribuicoes="2",
                cod_inc_trib="1",
                escritorio_id=escritorio.id,
            )
        )
        s.commit()
        from src.documentos import ImportadorDeDocumentos

        ImportadorDeDocumentos(s, escritorio_id=escritorio.id).importar(nfe_xml())
        s.commit()
        yield s


def _documento(sessao) -> DocumentoFiscal:
    return sessao.execute(select(DocumentoFiscal)).scalars().one()


class TestNaGravacao:
    def test_aplicar_ajuste_recusa_texto_que_nao_e_numero(self, sessao):
        """`aplicar_ajuste` é por onde toda escrita passa: é ali a última trava."""
        documento = _documento(sessao)
        with pytest.raises(ValueError, match="valor_icms"):
            aplicar_ajuste(
                sessao,
                documento=documento,
                item=documento.itens[0],
                campo="valor_icms",
                valor_novo="abc",
                origem=ORIGEM_USUARIO,
            )
        assert not sessao.execute(select(AjusteFiscal)).scalars().all()

    def test_ajuste_gravado_com_virgula_antes_da_correcao_ainda_gera(self, sessao):
        """O que o defeito já gravou precisa continuar gerando o arquivo."""
        documento = _documento(sessao)
        item = documento.itens[0]
        sessao.add(
            AjusteFiscal(
                documento_id=documento.id,
                item_id=item.id,
                campo="valor_icms",
                valor_anterior="180.0",
                valor_novo="190,00",
                origem=ORIGEM_USUARIO,
            )
        )
        sessao.commit()

        assert valor_efetivo(item, "valor_icms", item_ajustes(sessao)) == 190.0
        empresa = sessao.get(Empresa, 1)
        GeradorEFDICMS(sessao, empresa=empresa, data_inicio=INICIO, data_fim=FIM).gerar()
        GeradorEFDContribuicoes(sessao, empresa=empresa, data_inicio=INICIO, data_fim=FIM).gerar()

    def test_acao_de_regra_com_texto_invalido_e_recusada_no_cadastro(self, sessao):
        with pytest.raises(RegraInvalida, match="valor_icms"):
            criar_regra(
                sessao,
                nome="icms-errado",
                condicoes=[{"campo": "ncm", "valor": "22030000"}],
                acoes=[{"campo": "valor_icms", "valor": "abc"}],
            )

    def test_acao_de_regra_com_virgula_sugere_numero(self, sessao):
        """Sugestão em texto teria impacto `None` e compararia errado com o atual."""
        criar_regra(
            sessao,
            nome="icms-190",
            condicoes=[{"campo": "ncm", "valor": "22030000"}],
            acoes=[{"campo": "valor_icms", "valor": "190,00"}],
        )
        (sugestao,) = MotorDeClassificacao(sessao).avaliar(_documento(sessao)).sugestoes

        assert sugestao.valor_sugerido == 190.0
        assert sugestao.impacto == pytest.approx(10.0), (
            "a sugestão da regra não tem impacto em reais: a simulação esconderia a "
            "diferença que decide se ela passa"
        )

    def test_regra_que_repete_o_valor_atual_nao_sugere_nada(self, sessao):
        """ "180" e 180.0 são o mesmo ICMS — comparar como texto os separava."""
        criar_regra(
            sessao,
            nome="icms-180",
            condicoes=[{"campo": "ncm", "valor": "22030000"}],
            acoes=[{"campo": "valor_icms", "valor": "180"}],
        )
        assert not MotorDeClassificacao(sessao).avaliar(_documento(sessao)).sugestoes

    def test_planilha_com_texto_invalido_recusa_a_linha(self, sessao):
        conteudo = exportar(sessao, Selecao(escritorio_id=1, empresa_id=1))
        livro = load_workbook(io.BytesIO(conteudo))
        aba = livro["itens"]
        cabecalho = [c.value for c in aba[1]]
        aba.cell(row=2, column=cabecalho.index("valor_icms") + 1).value = "cento e noventa"
        aba.cell(row=2, column=cabecalho.index("cfop") + 1).value = "2102"
        buffer = io.BytesIO()
        livro.save(buffer)

        resultado = reimportar(sessao, buffer.getvalue())

        assert (
            not resultado.simulacao.mudancas
        ), "a linha com valor inválido virou alteração: o texto chegaria ao banco"
        (divergencia,) = resultado.divergencias
        assert divergencia.linha == 2 and "valor_icms" in divergencia.motivo

    def test_planilha_com_virgula_vira_numero(self, sessao):
        conteudo = exportar(sessao, Selecao(escritorio_id=1, empresa_id=1))
        livro = load_workbook(io.BytesIO(conteudo))
        aba = livro["itens"]
        cabecalho = [c.value for c in aba[1]]
        aba.cell(row=2, column=cabecalho.index("valor_icms") + 1).value = "190,00"
        buffer = io.BytesIO()
        livro.save(buffer)

        mudancas = reimportar(sessao, buffer.getvalue()).simulacao.mudancas

        (do_item,) = [m for m in mudancas if m.item_id is not None]
        assert do_item.valor_novo == 190.0


def item_ajustes(sessao):
    return sessao.execute(select(AjusteFiscal)).scalars().all()
