"""Demonstrações de uma ECD real: balanço, DRE, DFC, índices, assinaturas e plano.

Uma ECD de 2025 validada pelo PVA saiu errada em todos os relatórios: o
balanço sem saldo anterior e com "Total do Patrimônio Líquido 0,00", a DRE
com tudo em "Despesas Operacionais" (custo zero), o bloco FILTROS vazio no
PDF e a DFC contando a PECLD como caixa e as transferências entre contas da
própria empresa como fluxo. `tests/fixtures/ecd_demonstracoes.py` reproduz
cada defeito com números pequenos; os valores esperados abaixo são contas de
cabeça sobre os lançamentos de lá.

    Abertura: caixa 500, banco A 5.000, banco B 1.000, aplicação 2.000
    (caixa e equivalentes 8.500); clientes 3.000, PECLD (100), estoque 4.000,
    investimentos 1.500, veículos 10.000, depreciação (2.000) → ativo 24.900.
    Fornecedores 3.000, empréstimo 4.000, financiamento 5.000, capital
    10.000, lucros 2.900.

    Lucro 2.640 = receita 10.000 − ICMS 1.200 − compras 5.000 + devolução
    500 − energia 300 − salários 1.000 − depreciação 600 + outras receitas
    250 + rendimento 40 − IOF 20 − tarifas 30.

    Caixa: +8.000 clientes, −4.000 fornecedor, −300 energia, −1.200 ICMS,
    −1.000 salários, −50 IOF e tarifas, +250 receita eventual, +40
    rendimento → operacional 1.740; −4.000 veículo, −700 mútuo → investimento
    −4.700; +6.000 empréstimo, −1.500 amortização → financiamento 4.500.
    Variação 1.540 = 10.040 − 8.500. Fora: 4 transferências entre contas da
    empresa (3.000 + 2.000 + 2.000 + 1.000 = 8.000), o veículo financiado
    (2.000), a depreciação, a devolução abatida do fornecedor.
"""

from __future__ import annotations

from pathlib import Path
from unittest import mock

import pytest
from sqlalchemy import select

import src.ecd_importer as importador
from src.db.models import Empresa, Signatario, criar_engine, get_session, init_db
from src.ecd_importer import ECDImportService
from src.filters.engine import FilterCriteria, FilterEngine
from src.reports import documentos
from src.reports.assinaturas import Responsavel, assinaturas
from src.reports.balanco import BalancoPatrimonial
from src.reports.base import ReportContext
from src.reports.classificacao import Classificador
from src.reports.dfc import DFC
from src.reports.dre import DRE
from src.reports.indices import IndicesFinanceiros
from src.reports.plano_contas import PlanoDeContas
from src.validators.integridade import ValidadorIntegridade
from tests.fixtures.ecd_demonstracoes import (
    CONTADORA,
    CPF_CONTADORA,
    gerar_ecd_demonstracoes,
)


@pytest.fixture
def sessao():
    engine = criar_engine(":memory:")
    init_db(engine)
    s = get_session(engine)
    yield s
    s.close()
    engine.dispose()


def _importar(sessao, arquivo: Path) -> int:
    with mock.patch.object(importador, "emitir", lambda *a, **k: None):
        return ECDImportService(sessao).importar(arquivo).ecd_id


@pytest.fixture
def ecd(sessao, tmp_path) -> int:
    return _importar(sessao, gerar_ecd_demonstracoes(tmp_path / "ecd.txt"))


def _valores(linhas, campo="valor") -> dict[str, float]:
    return {
        ln.descricao: getattr(ln, campo) for ln in linhas if ln.tipo not in ("section", "grupo")
    }


# ── Importação: J930 ───────────────────────────────────────────────────────


class TestSignatarios:
    def test_j930_entra_com_o_cpf_intacto(self, sessao, ecd):
        signatarios = list(
            sessao.execute(select(Signatario).where(Signatario.ecd_id == ecd)).scalars()
        )
        contador = next(s for s in signatarios if s.cod_assin == "900")
        assert contador.nome == CONTADORA
        assert contador.cpf_cnpj == CPF_CONTADORA, "lido como número, o CPF perdia o zero"
        assert (contador.crc, contador.uf_crc) == ("012345/O-6", "SP")
        empresa = next(s for s in signatarios if s.ind_resp_legal == "S")
        assert empresa.cod_assin == "001" and len(empresa.cpf_cnpj) == 14

    def test_email_e_telefone_nao_sao_guardados(self):
        assert not {"email", "fone"} & set(Signatario.__table__.columns.keys())


# ── Estrutura: superior do I050 × balanço publicado ────────────────────────


class TestEstrutura:
    def test_superior_trocado_segue_o_j100(self, sessao, ecd):
        correcoes = FilterEngine(sessao, ecd).correcoes_de_superior()
        assert correcoes == {
            "1.1.2.01": ("1.1.1.03", "1.1.2"),
            "1.1.2.02": ("1.1.1.03", "1.1.2"),
            "1.2.2.01": ("1.2.1.01", "1.2.2"),
        }
        hierarquia = FilterEngine(sessao, ecd).hierarquia()
        assert hierarquia.superior("1.1.2.01") == "1.1.2"

    def test_sem_j100_nao_corrige_nada(self, sessao, tmp_path):
        ecd = _importar(sessao, gerar_ecd_demonstracoes(tmp_path / "s.txt", com_j100=False))
        assert FilterEngine(sessao, ecd).correcoes_de_superior() == {}

    def test_validacao_avisa_cada_correcao(self, sessao, ecd):
        alertas = [
            i
            for i in ValidadorIntegridade(sessao, ecd).validar_todas()
            if i.tipo == "superior_divergente_do_publicado"
        ]
        assert sorted(i.detalhes["cod_cta"] for i in alertas) == [
            "1.1.2.01",
            "1.1.2.02",
            "1.2.2.01",
        ]
        assert all(i.severidade == "alerta" for i in alertas)

    def test_conta_de_resultado_fora_da_dre_publicada(self, sessao, tmp_path):
        """A DRE publicada que omite contas: nenhuma linha diverge, o total sim."""
        texto = gerar_ecd_demonstracoes(tmp_path / "b.txt").read_text(encoding="utf-8")
        sem_devolucao = "\n".join(
            ln for ln in texto.splitlines() if not ln.startswith("|J150|") or "|4.1.1.02|" not in ln
        )
        arquivo = tmp_path / "c.txt"
        arquivo.write_text(sem_devolucao + "\n", encoding="utf-8")
        ecd = _importar(sessao, arquivo)
        alertas = [
            i
            for i in ValidadorIntegridade(sessao, ecd).validar_todas()
            if i.tipo == "conta_fora_da_dre_publicada"
        ]
        assert [i.detalhes["cod_cta"] for i in alertas] == ["4.1.1.02"]


# ── Balanço ────────────────────────────────────────────────────────────────


class TestBalanco:
    def test_saldo_anterior_e_o_de_abertura(self, sessao, ecd):
        """Sem ECD anterior importada, a coluna saía toda zerada."""
        _ctx, grupos, totais = BalancoPatrimonial(sessao, ecd).gerar()
        ativo = {ln.cod_cta: ln for ln in grupos["ativo"]}
        assert (ativo["1"].saldo_atual, ativo["1"].saldo_anterior) == (34_540, 24_900)
        assert ativo["1.1.1"].saldo_anterior == 8_500
        assert totais["tem_anterior"] is True and totais["origem_anterior"] == "saldo_inicial"
        assert (totais["ativo_anterior"], totais["passivo_anterior"], totais["pl_anterior"]) == (
            24_900,
            12_000,
            12_900,
        )
        assert str(totais["data_anterior"]) == "2024-12-31"

    def test_pl_de_natureza_02_vai_para_o_pl(self, sessao, ecd):
        _ctx, grupos, totais = BalancoPatrimonial(sessao, ecd).gerar()
        assert (totais["passivo"], totais["pl"], totais["passivo_pl"]) == (19_000, 15_540, 34_540)
        assert {ln.cod_cta for ln in grupos["pl"]} == {"2.3", "2.3.1", "2.3.2"}
        assert "2" not in {
            ln.cod_cta for ln in grupos["passivo"]
        }, "a sintética PASSIVO soma passivo e PL: não é nem um nem outro, é o total"

    def test_disponivel_sem_os_clientes(self, sessao, ecd):
        _ctx, grupos, _totais = BalancoPatrimonial(sessao, ecd).gerar()
        ativo = {ln.cod_cta: ln.saldo_atual for ln in grupos["ativo"]}
        assert ativo["1.1.1"] == 10_040, "com o I050 à risca, clientes e PECLD somavam aqui"
        assert ativo["1.1.2"] == 4_900
        assert ativo["1.2.1"] == 700 and ativo["1.2.2"] == 1_500

    def test_visao_de_publicacao_e_o_j100(self, sessao, ecd):
        ctx, grupos, totais = BalancoPatrimonial(sessao, ecd).gerar_publicacao()
        assert ctx.titulo == "Balanço Patrimonial (Publicação)"
        assert (totais["ativo"], totais["passivo"], totais["pl"]) == (34_540, 19_000, 15_540)
        assert (totais["ativo_anterior"], totais["pl_anterior"]) == (24_900, 12_900)
        assert totais["origem_anterior"] == "publicado"
        assert {ln.cod_cta for ln in grupos["pl"]} == {"2.3", "2.3.1", "2.3.2"}


# ── DRE ────────────────────────────────────────────────────────────────────


class TestDRE:
    def test_cada_degrau(self, sessao, ecd):
        _ctx, linhas, totais = DRE(sessao, ecd).gerar()
        valores = {ln.descricao: ln.valor_atual for ln in linhas}
        assert valores == {
            "Receita Operacional Bruta": 10_000,
            "(-) Deduções da Receita": -1_200,
            "= Receita Operacional Líquida": 8_800,
            "(-) Custos": -4_500,
            "= Lucro Bruto": 4_300,
            "(-) Despesas Operacionais": -1_900,
            "(+/-) Outras Receitas e Despesas Operacionais": 250,
            "= Resultado Antes do Resultado Financeiro": 2_650,
            "(+) Receitas Financeiras": 40,
            "(-) Despesas Financeiras": -50,
            "= Resultado Antes IRPJ/CSLL": 2_640,
            "(-) IRPJ / CSLL": 0,
            "= Resultado Líquido do Exercício": 2_640,
        }
        assert totais["resultado_liquido"] == 2_640

    def test_classificacao_pelo_referencial_e_pelo_grupo(self, sessao, ecd):
        classificador = Classificador(FilterEngine(sessao, ecd))
        assert classificador.categoria_dre("4.1.1.01") == "custos", "COMPRAS DE MERCADORIAS"
        assert (
            classificador.categoria_dre("4.1.1.02") == "custos"
        ), "(-) DEVOLUÇÃO DE MERCADORIAS dentro do CMV é devolução de compra, não dedução"
        assert classificador.categoria_dre("4.3.01") == "despesas_financeiras", "IOF"
        assert classificador.categoria_dre("4.3.02") == "despesas_financeiras", "TARIFAS"
        assert classificador.categoria_dre("4.4.01") == "receitas_financeiras"
        assert classificador.categoria_dre("4.4.02") == "outras_operacionais"
        assert classificador.categoria_dre("5.1") is None, "apuração (natureza 09) fica fora"

    def test_anterior_vem_da_dre_publicada(self, sessao, ecd):
        _ctx, linhas, totais = DRE(sessao, ecd).gerar()
        anterior = {ln.descricao: ln.valor_anterior for ln in linhas}
        assert anterior["Receita Operacional Bruta"] == 8_000
        assert anterior["(-) Custos"] == -4_000
        assert anterior["= Resultado Líquido do Exercício"] == 2_840
        assert totais["origem_anterior"] == "publicado"

    def test_detalhe_por_conta(self, sessao, ecd):
        _ctx, linhas, _totais = DRE(sessao, ecd).gerar(detalhar=True)
        detalhes = [ln for ln in linhas if ln.tipo == "detail"]
        custos = [ln.cod_cta for ln in detalhes if ln.ordem == 3]
        assert custos == ["4.1.1.01", "4.1.1.02"]
        assert next(ln for ln in detalhes if ln.cod_cta == "4.1.1.02").valor_atual == 500


# ── DFC ────────────────────────────────────────────────────────────────────


DIRETO = {
    "Recebimentos de clientes": 8_000,
    "Juros e rendimentos recebidos": 40,
    "Outros recebimentos operacionais": 250,
    "Pagamentos a fornecedores de mercadorias e serviços": -4_300,
    "Pagamentos a empregados e encargos sociais": -1_000,
    "Pagamentos de tributos (exceto IR e CSLL)": -1_200,
    "Juros e encargos financeiros pagos": -50,
    "= Caixa líquido das atividades operacionais": 1_740,
    "Aquisição de imobilizado": -4_000,
    "Empréstimos concedidos a terceiros": -700,
    "= Caixa líquido das atividades de investimento": -4_700,
    "Captação de empréstimos e financiamentos": 6_000,
    "Amortização de empréstimos e financiamentos": -1_500,
    "= Caixa líquido das atividades de financiamento": 4_500,
    "= Aumento (redução) líquido de caixa e equivalentes": 1_540,
    "Caixa e equivalentes no início do período": 8_500,
    "Caixa e equivalentes no fim do período": 10_040,
    "Variação de caixa nos saldos (fim − início)": 1_540,
    "Diferença não conciliada (fluxos − saldos)": 0,
}

INDIRETO_OPERACIONAL = {
    "Lucro (prejuízo) líquido do período": 2_640,
    "(+) Depreciação, amortização e exaustão": 600,
    "(Aumento) redução em clientes": -2_000,
    "Aumento (redução) em fornecedores": 500,
    "= Caixa líquido das atividades operacionais": 1_740,
}


class TestDFC:
    def test_metodo_direto(self, sessao, ecd):
        _ctx, linhas, totais = DFC(sessao, ecd).gerar(metodo="direto")
        assert _valores(linhas) == DIRETO
        assert totais["conciliado"] is True

    def test_metodo_indireto(self, sessao, ecd):
        _ctx, linhas, totais = DFC(sessao, ecd).gerar(metodo="indireto")
        valores = _valores(linhas)
        assert {k: valores.get(k) for k in INDIRETO_OPERACIONAL} == INDIRETO_OPERACIONAL
        assert totais["variacao_caixa"] == 1_540 and totais["conciliado"] is True

    def test_os_dois_metodos_dao_o_mesmo_caixa_das_operacoes(self, sessao, ecd):
        _ctx, _linhas, totais = DFC(sessao, ecd).gerar_ambos()
        assert totais["operacional_direto"] == totais["operacional_indireto"] == 1_740
        assert totais["diferenca_metodos"] == 0

    def test_transferencia_entre_contas_da_empresa_nao_e_fluxo(self, sessao, ecd):
        """Banco → banco, por conta de passagem, e banco → aplicação: 4 lançamentos, 8.000."""
        _ctx, _linhas, totais = DFC(sessao, ecd).gerar()
        assert totais["transferencias_internas_qtd"] == 4
        assert totais["transferencias_internas"] == 8_000
        caixa = {c["cod_cta"] for c in totais["composicao_caixa"]}
        assert "2.1.2.02" in caixa, "a conta de passagem é numerário em trânsito"

    def test_transacao_sem_caixa_fica_fora(self, sessao, ecd):
        """O veículo financiado (2.000) não é aquisição nem captação."""
        _ctx, linhas, _totais = DFC(sessao, ecd).gerar(metodo="direto")
        valores = _valores(linhas)
        assert valores["Aquisição de imobilizado"] == -4_000
        assert valores["Captação de empréstimos e financiamentos"] == 6_000

    def test_pecld_pendurada_em_aplicacoes_nao_e_caixa(self, sessao, tmp_path):
        """Sem o J100 para corrigir o superior, o nome da própria conta decide."""
        ecd = _importar(sessao, gerar_ecd_demonstracoes(tmp_path / "s.txt", com_j100=False))
        classificador = Classificador(FilterEngine(sessao, ecd))
        assert not classificador.eh_caixa("1.1.2.02"), "(-) PERDAS ESTIMADAS…"
        assert not classificador.eh_caixa("1.1.2.01"), "CLIENTES A RECEBER"
        assert classificador.eh_caixa("1.1.1.03.01")
        _ctx, _linhas, totais = DFC(sessao, ecd).gerar()
        assert (totais["caixa_inicial"], totais["caixa_final"]) == (8_500, 10_040)

    def test_periodo_restringe_os_lancamentos(self, sessao, ecd):
        import datetime

        criterios = FilterCriteria(
            dt_ini=datetime.date(2025, 1, 1), dt_fin=datetime.date(2025, 1, 31)
        )
        _ctx, linhas, _totais = DFC(sessao, ecd).gerar(criterios, metodo="direto")
        assert _valores(linhas)["Recebimentos de clientes"] == 8_000
        assert "Aquisição de imobilizado" not in _valores(linhas)


# ── Índices para licitação ─────────────────────────────────────────────────


class TestIndices:
    def test_indices_usuais_e_anteriores(self, sessao, ecd):
        _ctx, indices, totais = IndicesFinanceiros(sessao, ecd).gerar()
        por_chave = {i.chave: i for i in indices}
        # AC 18.940, RLP 700, AT 34.540, PC 12.000, PNC 7.000, PL 15.540
        assert por_chave["lg"].valor == pytest.approx(19_640 / 19_000, abs=1e-4)
        assert por_chave["sg"].valor == pytest.approx(34_540 / 19_000, abs=1e-4)
        assert por_chave["lc"].valor == pytest.approx(18_940 / 12_000, abs=1e-4)
        assert por_chave["ls"].valor == pytest.approx(14_940 / 12_000, abs=1e-4)
        assert por_chave["li"].valor == pytest.approx(10_040 / 12_000, abs=1e-4)
        assert por_chave["ccl"].valor == 6_940
        # Abertura: AC 15.400, AT 24.900, PC 7.000, PNC 5.000
        assert por_chave["lc"].valor_anterior == pytest.approx(2.2)
        assert por_chave["sg"].valor_anterior == pytest.approx(2.075)
        assert all(i.atende for i in indices if i.usual)
        assert totais["atende_usuais"] is True
        assert totais["contrato_maximo_pl_10"] == 155_400

    def test_nao_atende_quando_o_indice_e_menor_que_1(self):
        from src.reports.indices import Indice

        indice = Indice("lc", "LC", "", 0.21, 1.54, referencia="> 1,00", usual=True)
        assert (indice.atende, indice.atende_anterior) == (False, True)


# ── Assinaturas ────────────────────────────────────────────────────────────


class TestAssinaturas:
    def test_contador_do_j930_e_socio_em_branco(self, sessao, ecd):
        responsavel, contador = assinaturas(sessao, ecd)
        assert contador.nome == CONTADORA
        assert "CRC-SP 012345/O-6" in contador.documento
        assert "CPF 012.345.678-90" in contador.documento
        assert responsavel.nome == "", "o e-CNPJ não diz quem é o sócio"
        assert "Sócio administrador" in responsavel.cargo

    def test_socio_informado_ou_cadastrado(self, sessao, ecd):
        informado = Responsavel("FULANO DE TAL", "11144477735", "Sócio administrador")
        assert assinaturas(sessao, ecd, informado)[0].documento == "CPF 111.444.777-35"

        empresa = sessao.execute(select(Empresa)).scalar_one()
        empresa.responsavel_nome = "BELTRANA SOCIA"
        empresa.responsavel_qualificacao = "Administradora"
        sessao.commit()
        responsavel = assinaturas(sessao, ecd)[0]
        assert (responsavel.nome, responsavel.cargo) == ("BELTRANA SOCIA", "Administradora")


# ── Plano de contas ────────────────────────────────────────────────────────


class TestPlanoDeContas:
    def test_extrai_o_plano_como_declarado(self, sessao, ecd):
        _ctx, linhas, totais = PlanoDeContas(sessao, ecd).gerar()
        assert totais["total"] == 58
        por_codigo = {ln.cod_cta: ln for ln in linhas}
        assert por_codigo["1.1.1.02.01"].referencial == "1.01.01.02.01"
        assert por_codigo["1.1.1.02.01"].aglutinacao == "1.1.1.02.01"
        assert por_codigo["1.1.2.01"].cod_cta_sup == "1.1.1.03", "o I050 como está"
        assert por_codigo["1.1.2.01"].superior_publicado == "1.1.2"
        assert totais["divergencias_publicado"] == 3
        assert totais["por_natureza"] == {
            "Ativo": 23,
            "Passivo": 13,
            "Resultado": 20,
            "Outras": 2,
        }

    def test_filtro_de_natureza(self, sessao, ecd):
        _ctx, linhas, _totais = PlanoDeContas(sessao, ecd).gerar(FilterCriteria(cod_nat=["04"]))
        assert linhas and all(ln.cod_nat == "04" for ln in linhas)


# ── Documentos: PDF, XLSX e TXT ────────────────────────────────────────────


class TestDocumentos:
    def test_filtros_so_quando_ha_filtro(self, sessao, ecd):
        assert not ReportContext(titulo="x").tem_filtros, "o painel montava o contexto vazio"
        assert not ReportContext(titulo="x", filtros_descricao="Nenhum filtro aplicado").tem_filtros
        assert ReportContext(titulo="x", filtros_descricao="Natureza: Ativo").tem_filtros

        sem = documentos.html(documentos.montar(sessao, ecd, "balanco"))
        assert "<dt>Filtros</dt>" not in sem
        com = documentos.html(
            documentos.montar(sessao, ecd, "balanco", FilterCriteria(cod_nat=["01"]))
        )
        assert "<dt>Filtros</dt>" in com and "Natureza: Ativo" in com

    @pytest.mark.parametrize("tipo", ["balanco", "dre", "dfc", "indices", "balancete"])
    def test_demonstracoes_saem_assinadas(self, sessao, ecd, tipo):
        html = documentos.html(documentos.montar(sessao, ecd, tipo))
        assert CONTADORA in html and "CRC-SP 012345/O-6" in html
        assert "Sócio administrador / Responsável legal" in html

    def test_plano_nao_leva_assinatura(self, sessao, ecd):
        html = documentos.html(documentos.montar(sessao, ecd, "plano"))
        assert 'class="assinaturas"' not in html and CONTADORA not in html

    def test_dfc_em_pdf_traz_os_dois_metodos(self, sessao, ecd):
        html = documentos.html(documentos.montar(sessao, ecd, "dfc"))
        assert "Método direto" in html and "Método indireto" in html
        assert "Recebimentos de clientes" in html and "(Aumento) redução em clientes" in html
        assert "4 lançamento(s)" in html, "o critério das transferências internas"

    def test_plano_em_texto(self, sessao, ecd):
        texto = documentos.texto(documentos.montar(sessao, ecd, "plano"))
        linhas = texto.splitlines()
        assert linhas[0] == "PLANO DE CONTAS"
        assert "11.222.333/0001-81" in linhas[1]
        assert any(ln.strip().startswith("1.1.1.02.01") and "BANCO ALFA" in ln for ln in linhas)
        assert "Contas: 58" in texto

    def test_pdf_e_pdf(self, sessao, ecd):
        conteudo, media_type, extensao = documentos.exportar(
            documentos.montar(sessao, ecd, "indices"), "pdf"
        )
        assert conteudo.startswith(b"%PDF") and media_type == "application/pdf"
        assert extensao == "pdf"

    def test_indices_so_com_os_valores(self, sessao, ecd):
        documento = documentos.montar(sessao, ecd, "indices")
        assert documento.colunas == ["indice", "formula", "valor", "valor_anterior"]
        html = documentos.html(documento)
        assert "Liquidez Geral (LG)" in html
        for removido in ("Situação", "Referência", "Não atende", "status-ok"):
            assert f">{removido}" not in html and f'"{removido}"' not in html, removido
        texto = documentos.texto(documento)
        assert "Atende" not in texto and "> 1,00" not in texto

    def test_xlsx_com_rotulos_e_datas(self, sessao, ecd):
        import io

        from openpyxl import load_workbook

        conteudo, _mt, _ext = documentos.exportar(documentos.montar(sessao, ecd, "balanco"), "xlsx")
        planilha = load_workbook(io.BytesIO(conteudo)).active
        cabecalho = next(ln for ln in planilha.iter_rows(values_only=True) if ln[0] == "Seção")
        assert cabecalho == ("Seção", "Código", "Nome", "31/12/2025", "31/12/2024")

    def test_tipo_desconhecido(self, sessao, ecd):
        with pytest.raises(documentos.TipoDeRelatorioDesconhecido):
            documentos.montar(sessao, ecd, "razonete")


# ── Pela linha de comando e pela tela (§7.1) ───────────────────────────────


@pytest.fixture
def banco(tmp_path) -> str:
    from src.cli import main

    caminho = str(tmp_path / "demonstracoes.db")
    main(["importar-ecd", str(gerar_ecd_demonstracoes(tmp_path / "ecd.txt")), "--db", caminho])
    return caminho


class TestPelaCLI:
    def test_plano_de_contas_em_txt(self, banco, tmp_path):
        from src.cli import main

        saida = tmp_path / "plano.txt"
        main(["exportar", "plano", "--formato", "txt", "--saida", str(saida), "--db", banco])
        texto = saida.read_text(encoding="utf-8")
        assert texto.startswith("PLANO DE CONTAS\n")
        assert "TRANSFERÊNCIAS ENTRE CONTAS" in texto and "1.01.01.05.01" in texto

    def test_plano_de_contas_em_pdf(self, banco, tmp_path):
        from src.cli import main

        saida = tmp_path / "plano.pdf"
        main(["exportar", "plano", "--formato", "pdf", "--saida", str(saida), "--db", banco])
        assert saida.read_bytes().startswith(b"%PDF")

    def test_socio_na_assinatura(self, banco, tmp_path):
        from src.cli import main

        saida = tmp_path / "indices.txt"
        main(
            [
                "exportar",
                "indices",
                "--formato",
                "txt",
                "--saida",
                str(saida),
                "--socio",
                "FULANO DE TAL",
                "--socio-cpf",
                "11144477735",
                "--db",
                banco,
            ]
        )
        texto = saida.read_text(encoding="utf-8")
        assert "FULANO DE TAL" in texto and "CPF 111.444.777-35" in texto
        assert CONTADORA in texto and "Liquidez Geral (LG)" in texto

    def test_sem_assinaturas(self, banco, tmp_path):
        from src.cli import main

        saida = tmp_path / "balanco.txt"
        main(
            ["exportar", "balanco", "--formato", "txt", "--saida", str(saida)]
            + ["--sem-assinaturas", "--db", banco]
        )
        assert CONTADORA not in saida.read_text(encoding="utf-8")

    def test_relatorio_dfc_por_metodo(self, banco, capsys):
        from src.cli import main

        main(["relatorio", "dfc", "--metodo", "direto", "--db", banco])
        direto = capsys.readouterr().out
        assert "Recebimentos de clientes" in direto and "Lucro (prejuízo)" not in direto
        main(["relatorio", "dfc", "--metodo", "indireto", "--db", banco])
        indireto = capsys.readouterr().out
        assert "Lucro (prejuízo) líquido do período" in indireto and "1.740,00" in indireto

    def test_relatorio_balanco_mostra_o_anterior(self, banco, capsys):
        from src.cli import main

        main(["relatorio", "balanco", "--db", banco])
        saida = capsys.readouterr().out
        assert "31/12/2024" in saida and "24,900.00" in saida
        assert "-0.00" not in saida


class TestPelaTela:
    @pytest.fixture
    def cliente(self, banco, monkeypatch):
        from fastapi.testclient import TestClient

        from src.auth import init_auth

        monkeypatch.setenv("SPED_HUB_DB", banco)
        from src.dashboard.app import app

        init_auth(banco)
        cliente = TestClient(app)
        dados = {"email": "demo@test.local", "nome": "Demo", "senha": "senha123"}
        assert cliente.post("/api/register", data=dados).status_code == 200
        del dados["nome"]
        assert cliente.post("/api/login", data=dados).status_code == 200
        return cliente

    def test_abas_novas(self, cliente):
        painel = cliente.get("/").text
        assert "/api/indices?ecd_id=1" in painel and "/api/plano?ecd_id=1" in painel
        indices = cliente.get("/api/indices", params={"ecd_id": 1}).text
        assert "Liquidez Geral (LG)" in indices and "1,03" in indices
        assert (
            "Atende" not in indices and "Situação" not in indices
        ), "a tabela mostra só os índices: sem situação nem referência"
        assert "Referência" not in indices
        plano = cliente.get("/api/plano", params={"ecd_id": 1}).text
        assert "BANCO ALFA" in plano and "J100: sob 1.1.2" in plano

    def test_dfc_pelos_dois_metodos(self, cliente):
        direto = cliente.get("/api/dfc", params={"ecd_id": 1}).text
        assert "Recebimentos de clientes" in direto and "1.740,00" in direto
        indireto = cliente.get("/api/dfc", params={"ecd_id": 1, "metodo": "indireto"}).text
        assert "(Aumento) redução em clientes" in indireto and "1.740,00" in indireto
        assert cliente.get("/api/dfc", params={"ecd_id": 1, "metodo": "x"}).status_code == 422

    def test_balanco_com_saldo_anterior(self, cliente):
        html = cliente.get("/api/balanco", params={"ecd_id": 1}).text
        assert "31/12/2024" in html and "24.900,00" in html

    def test_assinatura_do_socio_pela_tela(self, cliente):
        formulario = cliente.get("/api/assinaturas", params={"ecd_id": 1}).text
        assert CONTADORA in formulario
        resposta = cliente.post(
            "/api/assinaturas?ecd_id=1",
            data={"nome": "FULANO DE TAL", "cpf": "111.444.777-35", "qualificacao": ""},
        )
        assert resposta.status_code == 200 and "Responsável salvo" in resposta.text
        texto = cliente.get("/api/export/txt", params={"ecd_id": 1, "tipo": "balanco"}).text
        assert "FULANO DE TAL" in texto and "Sócio administrador" in texto
        invalido = cliente.post("/api/assinaturas?ecd_id=1", data={"nome": "X", "cpf": "123"})
        assert invalido.status_code == 400

    @pytest.mark.parametrize("formato", ["pdf", "xlsx", "txt"])
    def test_exportacoes(self, cliente, formato):
        for tipo in ("plano", "indices", "dfc"):
            resposta = cliente.get(f"/api/export/{formato}", params={"ecd_id": 1, "tipo": tipo})
            assert resposta.status_code == 200, (tipo, resposta.text[:200])
            assert f"{tipo}_1.{formato}" in resposta.headers["content-disposition"]

    def test_pdf_do_painel_sem_bloco_de_filtros(self, cliente, monkeypatch):
        """O painel montava o contexto sem descrição: o PDF saía com FILTROS vazio."""
        capturado = {}
        original = documentos.pdf

        def espiao(documento, white_label=None):
            capturado["html"] = documentos.html(documento, white_label)
            return original(documento, white_label)

        monkeypatch.setattr(documentos, "pdf", espiao)
        assert (
            cliente.get("/api/export/pdf", params={"ecd_id": 1, "tipo": "dre"}).status_code == 200
        )
        assert "<dt>Filtros</dt>" not in capturado["html"]
        assert "01/01/2025 a 31/12/2025" in capturado["html"], "período em pt-BR, não ISO"


# ── Repartição do caixa entre as contrapartidas ────────────────────────────


class TestRateio:
    def test_so_o_lado_oposto_ao_caixa_recebe(self):
        from src.reports.dfc import ratear

        # Recebimento de 90 com desconto de 10: C clientes 100, D desconto 10.
        assert ratear(-90, [10, -100]) == pytest.approx([0, -90])
        assert ratear(-50, [-20, -30]) == pytest.approx([-20, -30])
        assert ratear(0, [5, -5]) == [0, 0]

    def test_recebimento_com_desconto_concedido(self, sessao, tmp_path):
        """Entraram 90: é recebimento de clientes de 90, e o desconto não é caixa.

        Com o valor de cada contrapartida (em vez da repartição do caixa pelo
        lado oposto), saíam recebimento de 100 e "juros pagos" de 10 — dois
        fluxos que não aconteceram.
        """
        linhas = [
            "|0000|LECD|01012025|31122025|DESCONTO LTDA|11222333000181|SP||3550308"
            "|||0|1|0||1|0||N|N|0|0||",
            "|I001|0|",
            "|I010|G|009|",
            "|I050|01012025|01|S|1|1||ATIVO|",
            "|I050|01012025|01|A|2|1.1|1|CAIXA|",
            "|I050|01012025|01|A|2|1.2|1|CLIENTES|",
            "|I050|01012025|02|S|1|2||PATRIMÔNIO LÍQUIDO|",
            "|I050|01012025|02|A|2|2.1|2|CAPITAL SOCIAL|",
            "|I050|01012025|02|A|2|2.2|2|LUCROS ACUMULADOS|",
            "|I050|01012025|04|S|1|3||RESULTADO|",
            "|I050|01012025|04|A|2|3.1|3|RECEITA DE VENDAS|",
            "|I050|01012025|04|A|2|3.2|3|DESCONTOS CONCEDIDOS|",
            "|I150|01012025|31122025|",
            "|I155|1.1||50,00|D|90,00|0,00|140,00|D|",
            "|I155|1.2||0,00|D|100,00|100,00|0,00|D|",
            "|I155|2.1||50,00|C|0,00|0,00|50,00|C|",
            "|I155|2.2||0,00|C|0,00|90,00|90,00|C|",
            "|I155|3.1||0,00|C|100,00|100,00|0,00|C|",
            "|I155|3.2||0,00|D|10,00|10,00|0,00|D|",
            "|I200|1|10012025|100,00|N||",
            "|I250|1.2||100,00|D|||VENDA A PRAZO||",
            "|I250|3.1||100,00|C|||||",
            "|I200|2|20012025|100,00|N||",
            "|I250|1.1||90,00|D|||RECEBIMENTO COM DESCONTO||",
            "|I250|3.2||10,00|D|||||",
            "|I250|1.2||100,00|C|||||",
            "|I200|3|31122025|100,00|E||",
            "|I250|3.1||100,00|D|||ENCERRAMENTO||",
            "|I250|3.2||10,00|C|||||",
            "|I250|2.2||90,00|C|||||",
            "|I350|31122025|",
            "|I355|3.1||100,00|C|",
            "|I355|3.2||10,00|D|",
            "|I990|0|",
            "|9999|0|",
        ]
        arquivo = tmp_path / "desconto.txt"
        arquivo.write_text("\n".join(linhas) + "\n", encoding="utf-8")
        ecd = _importar(sessao, arquivo)

        _ctx, direto, totais = DFC(sessao, ecd).gerar(metodo="direto")
        assert _valores(direto)["Recebimentos de clientes"] == 90
        assert "Juros e encargos financeiros pagos" not in _valores(direto)
        assert totais["operacional_indireto"] == totais["operacional_direto"] == 90
        assert totais["conciliado"] is True


# ── Balancete de verificação ───────────────────────────────────────────────


class TestBalancete:
    def test_saldo_com_indicador_de_natureza(self):
        from src.reports.base import fmt_saldo_dc

        assert fmt_saldo_dc(1234.5) == "1.234,50 D"
        assert fmt_saldo_dc(-1234.5) == "1.234,50 C"
        assert fmt_saldo_dc(-0.004) == "0,00"

    def test_pdf_mostra_d_e_c(self, sessao, ecd):
        html = documentos.html(documentos.montar(sessao, ecd, "balancete"))
        assert "10.040,00 D" in html, "disponível devedor"
        assert "15.540,00 C" in html, "patrimônio líquido credor"
        assert "(15.540,00)" not in html, "o sinal não diz se a conta está credora ou devedora"

    @pytest.mark.parametrize("tipo", ["balancete", "plano"])
    def test_relatorio_largo_sai_deitado(self, sessao, ecd, tipo):
        """Sete colunas num A4 em pé cortavam a última e quebravam os valores."""
        html = documentos.html(documentos.montar(sessao, ecd, tipo))
        assert '<body class="paisagem">' in html
        assert "size: A4 landscape" in html

    def test_demonstracao_segue_em_pe(self, sessao, ecd):
        assert '<body class="">' in documentos.html(documentos.montar(sessao, ecd, "balanco"))

    def test_texto_e_planilha_com_d_e_c(self, sessao, ecd):
        texto = documentos.texto(documentos.montar(sessao, ecd, "balancete"))
        assert "BALANCETE DE VERIFICAÇÃO" in texto
        assert "10.040,00 D" in texto and "15.540,00 C" in texto

    def test_aba_do_painel(self, banco, monkeypatch):
        from fastapi.testclient import TestClient

        from src.auth import init_auth

        monkeypatch.setenv("SPED_HUB_DB", banco)
        from src.dashboard.app import app

        init_auth(banco)
        cliente = TestClient(app)
        dados = {"email": "balancete@test.local", "nome": "B", "senha": "senha123"}
        assert cliente.post("/api/register", data=dados).status_code == 200
        del dados["nome"]
        assert cliente.post("/api/login", data=dados).status_code == 200

        painel = cliente.get("/").text
        assert "/api/balancete?ecd_id=1" in painel, "o balancete não tinha aba no painel"
        assert "/api/export/pdf?ecd_id=1&tipo=balancete" in painel
        html = cliente.get("/api/balancete", params={"ecd_id": 1}).text
        assert "BANCO ALFA" in html and "10.040,00 D" in html
        assert "Conferência SI + D − C = SF: OK" in html
