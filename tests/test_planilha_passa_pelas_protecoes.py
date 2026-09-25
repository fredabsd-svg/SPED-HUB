"""A volta da planilha passa pelas mesmas proteções da alteração em massa.

O defeito: `reimportar` montava as mudanças linha a linha e devolvia a
`Simulacao` sem chamar o `_verificar`, a trava de documento cancelado e o
`recalcular` que `simular` chama. A planilha era a única escrita do sistema
que pulava as três:

  * NCM de sete dígitos e CST "1" — o zero à esquerda que o Excel come quando
    a célula vira número — passavam, e o validador do Fisco recusaria;
  * o ICMS de um item mudava e o `VL_ICMS` do C100 continuava o de antes:
    360,00 no cabeçalho contra 430,00 nos C190, e o espelho dizia "ok";
  * a nota cancelada podia ser corrigida.

Os testes passam pela porta de entrada — `sped-hub fiscal planilha` — porque é
por ela que a planilha volta.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from src.cli import main
from src.db.models import (
    AjusteFiscal,
    DocumentoFiscal,
    Empresa,
    Escritorio,
    criar_engine,
    get_session,
    init_db,
)
from src.escrituracoes.leiaute import EFD_ICMS
from tests.fixtures_nfe import nfe_xml

CNPJ = "98765432000198"


@pytest.fixture
def banco(tmp_path):
    """Uma entrada de dois itens: 180,00 de ICMS em cada, 360,00 no C100."""
    url = f"sqlite:///{tmp_path / 'fiscal.db'}"
    engine = criar_engine(url=url)
    init_db(engine)
    with get_session(engine) as sessao:
        sessao.add(Escritorio(nome="Teste", slug="teste"))
        sessao.commit()
        sessao.add(
            Empresa(
                cnpj=CNPJ,
                nome="COMERCIO EXEMPLO LTDA",
                uf="TO",
                ie="293456789",
                cod_mun="1721000",
                ind_perfil="A",
                ind_ativ="1",
                ind_ativ_contribuicoes="2",
                cod_inc_trib="1",
                escritorio_id=1,
            )
        )
        sessao.commit()
    engine.dispose()
    xml = tmp_path / "nfe.xml"
    xml.write_bytes(nfe_xml(itens=2))
    assert main(["fiscal", "importar", str(xml), "--escritorio", "1", "--db", url]) == 0
    return url


def _exportar(banco, tmp_path):
    destino = tmp_path / "itens.xlsx"
    assert (
        main(["fiscal", "planilha", "--empresa", "1", "--saida", str(destino), "--db", banco]) == 0
    )
    return destino


def _editar(arquivo, **celulas):
    """Muda células da primeira linha de item, como quem edita no Excel."""
    livro = load_workbook(io.BytesIO(arquivo.read_bytes()))
    aba = livro["itens"]
    cabecalho = [c.value for c in aba[1]]
    for coluna, valor in celulas.items():
        aba.cell(row=2, column=cabecalho.index(coluna) + 1).value = valor
    buffer = io.BytesIO()
    livro.save(buffer)
    arquivo.write_bytes(buffer.getvalue())
    return arquivo


def _voltar(banco, arquivo, *extras):
    return main(["fiscal", "planilha", "--arquivo", str(arquivo), "--db", banco, *extras])


def _ajustes(banco):
    with get_session(criar_engine(url=banco)) as sessao:
        return sessao.execute(select(AjusteFiscal)).scalars().all()


class TestFormato:
    def test_ncm_sem_o_zero_a_esquerda_e_recusado(self, banco, tmp_path, capsys):
        planilha = _editar(_exportar(banco, tmp_path), ncm=2031100)

        codigo = _voltar(banco, planilha, "--confirmar")
        saida = capsys.readouterr().out

        assert codigo == 1, (
            "a planilha gravou NCM de sete dígitos: o C170/0200 sai com um código "
            "que o validador do Fisco recusa"
        )
        assert "NCM" in saida and "2031100" in saida
        assert not _ajustes(banco)

    def test_cst_de_um_digito_e_recusado(self, banco, tmp_path, capsys):
        planilha = _editar(_exportar(banco, tmp_path), cst_pis=1)

        codigo = _voltar(banco, planilha, "--confirmar")

        assert codigo == 1, (
            "CST '1' passou: é o '01' que o Excel transformou em número, e o "
            "campo CST_PIS do C170 tem dois dígitos"
        )
        assert "CST_PIS" in capsys.readouterr().out
        assert not _ajustes(banco)

    def test_sem_confirmar_o_aviso_aparece(self, banco, tmp_path, capsys):
        """Quem lê a volta precisa ver o problema antes de pedir para gravar."""
        planilha = _editar(_exportar(banco, tmp_path), ncm=2031100)

        assert _voltar(banco, planilha) == 0
        assert "IMPEDITIVO" in capsys.readouterr().out


class TestCabecalho:
    def test_o_c100_acompanha_o_icms_corrigido_na_planilha(self, banco, tmp_path):
        planilha = _editar(_exportar(banco, tmp_path), valor_icms=250.0)
        assert _voltar(banco, planilha, "--confirmar") == 0

        destino = tmp_path / "efd.txt"
        assert (
            main(
                [
                    "fiscal",
                    "gerar",
                    "--empresa",
                    "1",
                    "--de",
                    "2026-07-01",
                    "--ate",
                    "2026-07-31",
                    "--saida",
                    str(destino),
                    "--db",
                    banco,
                ]
            )
            == 0
        )
        linhas = destino.read_bytes().decode("latin-1").split("\r\n")
        c100 = next(ln for ln in linhas if ln.startswith("|C100|")).split("|")[2:-1]
        c190 = [ln.split("|")[2:-1] for ln in linhas if ln.startswith("|C190|")]
        do_cabecalho = c100[EFD_ICMS["C100"].index("VL_ICMS")]
        dos_c190 = sum(float(c[EFD_ICMS["C190"].index("VL_ICMS")].replace(",", ".")) for c in c190)

        assert do_cabecalho == "430,00" and dos_c190 == pytest.approx(430.0), (
            f"C100 com VL_ICMS {do_cabecalho} e C190 somando {dos_c190:.2f}: a planilha "
            "mudou o item sem recompor o cabeçalho, e o validador confere os dois"
        )


class TestCancelado:
    def test_nota_cancelada_nao_se_corrige_pela_planilha(self, banco, tmp_path):
        planilha = _exportar(banco, tmp_path)
        with get_session(criar_engine(url=banco)) as sessao:
            sessao.execute(select(DocumentoFiscal)).scalars().one().situacao = "cancelado"
            sessao.commit()
        _editar(planilha, cfop="2102")

        assert _voltar(banco, planilha, "--confirmar") == 1, (
            "a planilha alterou nota cancelada: é a mesma trava de `alterar`, e "
            "alterar documento que não existe mais gera arquivo que o Fisco rejeita"
        )
        assert not _ajustes(banco)
